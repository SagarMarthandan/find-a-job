from __future__ import annotations

import csv
import html
import io
import re
import zipfile
from html.parser import HTMLParser
from xml.etree import ElementTree


MAX_EXPANDED = 40 * 1024 * 1024


class _TextHTML(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.hidden = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style", "noscript"}:
            self.hidden += 1
        elif tag in {"p", "div", "br", "li", "tr", "h1", "h2", "h3"} and not self.hidden:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style", "noscript"} and self.hidden:
            self.hidden -= 1

    def handle_data(self, data: str) -> None:
        if not self.hidden:
            self.parts.append(data)


def _plain(data: bytes) -> str:
    return data.decode("utf-8-sig")


def _zip(data: bytes) -> zipfile.ZipFile:
    archive = zipfile.ZipFile(io.BytesIO(data))
    if sum(item.file_size for item in archive.infolist()) > MAX_EXPANDED:
        archive.close()
        raise ValueError("Office document expands beyond 40 MiB")
    return archive


def _extract_text(data: bytes, suffix: str) -> str:
    text = _plain(data)
    if suffix in {".csv", ".tsv"}:
        rows = csv.reader(io.StringIO(text), delimiter="\t" if suffix == ".tsv" else ",")
        return "\n".join(" | ".join(cell for cell in row) for row in rows)
    if suffix == ".xml":
        return "\n".join(part.strip() for part in ElementTree.fromstring(text).itertext() if part.strip())
    return text


def _extract_html(data: bytes) -> str:
    parser = _TextHTML()
    parser.feed(_plain(data))
    return "".join(parser.parts)


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    pages = [page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages]
    text = "\n".join(f"[Page {index + 1}]\n{value}" for index, value in enumerate(pages) if value.strip())
    if not text.strip():
        raise ValueError("No selectable PDF text; scanned or image-only PDFs require OCR before upload")
    return text


def _extract_docx(data: bytes) -> str:
    with _zip(data):
        pass
    from docx import Document

    document = Document(io.BytesIO(data))
    return "\n".join(
        [*(paragraph.text for paragraph in document.paragraphs),
         *(" | ".join(cell.text for row in table.rows for cell in row.cells) for table in document.tables)]
    )


def _extract_odt(data: bytes) -> str:
    with _zip(data) as archive:
        root = ElementTree.fromstring(archive.read("content.xml"))
    return "\n".join(part.strip() for part in root.itertext() if part.strip())


def _extract_xlsx(data: bytes) -> str:
    with _zip(data) as archive:
        archive.read("xl/workbook.xml")
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        return "\n".join(
            f"[{sheet.title}]\n" + "\n".join(
                " | ".join("" if cell is None else str(cell) for cell in row)
                for row in sheet.iter_rows(values_only=True)
            )
            for sheet in workbook.worksheets
        )
    finally:
        workbook.close()


def _extract_pptx(data: bytes) -> str:
    with _zip(data):
        pass
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    return "\n".join(
        f"[Slide {index + 1}]\n" + "\n".join(shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False))
        for index, slide in enumerate(presentation.slides)
    )


def _extract_rtf(data: bytes) -> str:
    from striprtf.striprtf import rtf_to_text

    return rtf_to_text(_plain(data))


_TEXT_FORMATS = {".txt", ".md", ".markdown", ".json", ".yaml", ".yml", ".csv", ".tsv", ".py", ".js", ".ts", ".java", ".sql", ".toml", ".ini", ".xml"}
_EXTRACTORS = {
    ".html": _extract_html,
    ".htm": _extract_html,
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".odt": _extract_odt,
    ".xlsx": _extract_xlsx,
    ".pptx": _extract_pptx,
    ".rtf": _extract_rtf,
}


def extract(data: bytes, suffix: str) -> str:
    suffix = suffix.lower()
    if suffix in _TEXT_FORMATS:
        text = _extract_text(data, suffix)
    elif extractor := _EXTRACTORS.get(suffix):
        text = extractor(data)
    else:
        raise ValueError(f"Unsupported file format {suffix or '(no extension)'}; export as PDF, DOCX, ODT, XLSX, PPTX, RTF, or UTF-8 text")
    text = html.unescape(re.sub(r"[\t\r ]+", " ", text))
    if not text.strip():
        raise ValueError("No readable text was extracted")
    if len(text) > 100_000:
        raise ValueError("Extracted text exceeds 100000 characters")
    return text
